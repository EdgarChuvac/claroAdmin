"""Pruebas HTTP de la API (backend en memoria, sin credenciales)."""

import io
import re

import openpyxl
from conftest import OPERATOR, SAMPLE_XLSX

OP_RE = re.compile(r"^OP-\d{8}-[0-9A-F]{12}$")


def _xlsx(build) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    build(wb, ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _small_inventory(ws, sheet="10.99.1.0", labels=None):
    """Subred /29 en columnas A/B: red .0, GW .1, hosts .2-.6, broadcast .7."""
    ws.title = sheet
    labels = labels or {}
    ws.append([0, "INTERNET 3001"])
    ws.append([1, "GW"])
    for octet in range(2, 7):
        ws.append([octet, labels.get(octet, "")])
    ws.append([7, "BROADCAST"])


def test_status_reports_backend_and_operation_id(client):
    res = client.get("/api/status")
    assert res.status_code == 200
    body = res.json()
    assert body["backend"] == "memory"
    assert OP_RE.match(res.headers["X-Operation-ID"])
    assert body["operation_id"] == res.headers["X-Operation-ID"]
    assert "/" not in str(body.get("excel_file", ""))


def test_lists_sheets_and_blocks(client):
    sheets = client.get("/api/sheets").json()["sheets"]
    assert "10.20.38.0" in sheets
    blocks = client.get("/api/blocks", params={"sheet": "10.20.38.0"}).json()["blocks"]
    first = blocks[0]
    assert first["network_ip"] == "10.20.38.0/27"
    assert first["gateway_ip"] == "10.20.38.1"
    assert first["first_available"]["ip"] == "10.20.38.6"
    assert first["assigned_count"] == 4


def test_unknown_sheet_returns_404_with_operation_id(client):
    res = client.get("/api/blocks", params={"sheet": "10.1.1.0"})
    assert res.status_code == 404
    assert OP_RE.match(res.json()["operation_id"])


def test_write_operations_require_operator(client):
    res = client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "S-1"})
    assert res.status_code == 400
    assert "operador" in res.json()["detail"]


def test_reserve_and_conflict(client):
    ok = client.post("/api/reservations", json={"ips": ["10.20.38.6", "10.20.38.7"], "service_id": "S-1"},
                     headers=OPERATOR)
    assert ok.status_code == 200, ok.text
    assert {r["ip"] for r in ok.json()["reserved"]} == {"10.20.38.6", "10.20.38.7"}

    # Atomicidad: si una IP falla, ninguna se reserva.
    again = client.post("/api/reservations", json={"ips": ["10.20.38.8", "10.20.38.7"], "service_id": "S-2"},
                        headers=OPERATOR)
    assert again.status_code == 409
    block = client.get("/api/blocks", params={"sheet": "10.20.38.0"}).json()["blocks"][0]
    assert "10.20.38.8" in {i["ip"] for i in block["available_ips"]}
    assert block["assigned_count"] == 6
    assigned = {i["ip"]: i for i in block["assigned_ips"]}
    assert assigned["10.20.38.6"]["assigned_by"] == "Operador Prueba"
    assert OP_RE.match(assigned["10.20.38.6"]["assigned_operation_id"])

    sheets = client.get("/api/sheets").json()["details"]
    sheet = next(s for s in sheets if s["sheet"] == "10.20.38.0")
    assert sheet["assigned_count"] == 6


def test_reserve_rejects_gateway_unknown_and_invalid(client):
    gw = client.post("/api/reservations", json={"ips": ["10.20.38.1"], "service_id": "S-1"}, headers=OPERATOR)
    assert gw.status_code == 409
    unknown = client.post("/api/reservations", json={"ips": ["10.250.0.5"], "service_id": "S-1"}, headers=OPERATOR)
    assert unknown.status_code == 404
    invalid = client.post("/api/reservations", json={"ips": ["999.1.1.1"], "service_id": "S-1"}, headers=OPERATOR)
    assert invalid.status_code == 400
    formula = client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "=CMD()"}, headers=OPERATOR)
    assert formula.status_code == 422
    blank = client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "   "}, headers=OPERATOR)
    assert blank.status_code == 422


def test_release_requires_matching_service(client):
    client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "S-1"}, headers=OPERATOR)
    wrong = client.post("/api/reservations/release",
                        json={"ips": ["10.20.38.6"], "service_id": "S-9", "reason": "baja"}, headers=OPERATOR)
    assert wrong.status_code == 409
    ok = client.post("/api/reservations/release",
                     json={"ips": ["10.20.38.6"], "service_id": "S-1", "reason": "baja del servicio"},
                     headers=OPERATOR)
    assert ok.status_code == 200
    block = client.get("/api/blocks", params={"sheet": "10.20.38.0"}).json()["blocks"][0]
    assert block["first_available"]["ip"] == "10.20.38.6"


def test_concurrent_reservations_only_one_wins(client):
    from concurrent.futures import ThreadPoolExecutor

    def attempt(n):
        return client.post("/api/reservations", json={"ips": ["10.20.38.9"], "service_id": f"S-{n}"},
                           headers=OPERATOR).status_code

    with ThreadPoolExecutor(max_workers=8) as pool:
        codes = list(pool.map(attempt, range(8)))
    assert codes.count(200) == 1
    assert codes.count(409) == 7


def test_import_merge_preserves_firestore_reservations(client):
    client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "S-1"}, headers=OPERATOR)
    res = client.post("/api/inventory/import", files={"file": ("inv.xlsx", SAMPLE_XLSX.read_bytes())},
                      data={"mode": "merge"}, headers=OPERATOR)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["conflicts_count"] == 1
    assert body["conflicts"][0]["ip"] == "10.20.38.6"
    ips = client.get("/api/services/S-1").json()["ips"]
    assert [i["ip"] for i in ips] == ["10.20.38.6"]


def test_import_replace_requires_confirmation_and_replaces(client):
    content = _xlsx(lambda wb, ws: _small_inventory(ws, labels={3: "SRV-X"}))
    no_confirm = client.post("/api/inventory/import", files={"file": ("a.xlsx", content)},
                             data={"mode": "replace"}, headers=OPERATOR)
    assert no_confirm.status_code == 400
    res = client.post("/api/inventory/import", files={"file": ("a.xlsx", content)},
                      data={"mode": "replace", "confirm_replace": "true"}, headers=OPERATOR)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["sheets"] == ["10.99.1.0"]
    assert body["totals"] == {"disponibles": 4, "ocupadas": 1, "gateways": 1}
    block = client.get("/api/blocks", params={"sheet": "10.99.1.0"}).json()["blocks"][0]
    assert block["network_ip"] == "10.99.1.0/29"
    assert [i["ip"] for i in block["available_ips"]] == ["10.99.1.2", "10.99.1.4", "10.99.1.5", "10.99.1.6"]
    # Las hojas no incluidas en el archivo se conservan.
    assert "10.20.38.0" in client.get("/api/sheets").json()["sheets"]


def test_import_reports_ignored_sheets_and_rejects_empty(client):
    def build(wb, ws):
        _small_inventory(ws)
        wb.create_sheet("Resumen")["A1"] = "no es inventario"

    res = client.post("/api/inventory/import", files={"file": ("a.xlsx", _xlsx(build))}, headers=OPERATOR)
    assert res.json()["ignored_sheets"] == ["Resumen"]

    empty = _xlsx(lambda wb, ws: setattr(ws, "title", "Hoja1"))
    res = client.post("/api/inventory/import", files={"file": ("b.xlsx", empty)}, headers=OPERATOR)
    assert res.status_code == 400


def test_import_rejects_bad_files(client):
    assert client.post("/api/inventory/import", files={"file": ("a.txt", b"hola")},
                       headers=OPERATOR).status_code == 400
    assert client.post("/api/inventory/import", files={"file": ("a.xlsx", b"PK-no-es-excel")},
                       headers=OPERATOR).status_code == 400


def test_export_roundtrip(client):
    client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "S-1"}, headers=OPERATOR)
    exported = client.get("/api/inventory/export")
    assert exported.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(exported.content))
    assert "10.20.38.0" in wb.sheetnames
    res = client.post("/api/inventory/import", files={"file": ("e.xlsx", exported.content)},
                      data={"mode": "merge"}, headers=OPERATOR)
    body = res.json()
    assert body["ips_written"] == 0 and body["conflicts_count"] == 0


def test_alta_is_recorded_and_idempotent(client):
    payload = {"data": {"id_servicio": "S-1", "cliente": "CLIENTE DEMO", "ip_wan": "10.20.38.6"}}
    first = client.post("/api/altas", json=payload, headers=OPERATOR)
    assert first.status_code == 200, first.text
    alta_id = first.json()["alta_id"]
    assert first.json()["duplicate"] is False
    second = client.post("/api/altas", json=payload, headers=OPERATOR)
    assert second.json() == {**second.json(), "alta_id": alta_id, "duplicate": True}

    doc = client.get(f"/api/altas/{alta_id}").json()
    assert doc["operator"] == "Operador Prueba"
    assert doc["service_id"] == "S-1"
    assert "ID DEL SERVICIO:  S-1" in doc["formatted_text"]
    assert OP_RE.match(doc["operation_id"])
    listed = client.get("/api/altas", params={"service_id": "S-1"}).json()["altas"]
    assert [a["alta_id"] for a in listed] == [alta_id]


def test_alta_requires_service_and_client(client):
    res = client.post("/api/altas", json={"data": {"id_servicio": "", "cliente": "X"}}, headers=OPERATOR)
    assert res.status_code == 422


def test_preview_does_not_record(client):
    res = client.post("/api/generate-format", json={"data": {"id_servicio": "S-1", "cliente": "X"}})
    assert res.status_code == 200
    assert client.get("/api/altas").json()["altas"] == []


def test_generate_format_rejects_unknown_fields(client):
    res = client.post("/api/generate-format", json={"data": {"campo_inexistente": "x"}})
    assert res.status_code == 422


def test_operations_are_audited(client):
    res = client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "S-1"}, headers=OPERATOR)
    op_id = res.headers["X-Operation-ID"]
    op = client.get(f"/api/operations/{op_id}").json()
    assert op["action"] == "ip.reserve"
    assert op["operator"] == "Operador Prueba"
    assert op["outcome"] == "ok"
    assert op["service_id"] == "S-1"
    assert op["details"]["ips"] == ["10.20.38.6"]
    assert op["expires_at"] > op["created_at"]

    failed = client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "S-2"}, headers=OPERATOR)
    fop = client.get(f"/api/operations/{failed.json()['operation_id']}").json()
    assert fop["outcome"] == "error" and fop["status_code"] == 409
    assert "ya no está disponible" in fop["error"]

    by_operator = client.get("/api/operations", params={"operator": "Operador Prueba"}).json()["operations"]
    assert {o["operation_id"] for o in by_operator} >= {op_id, failed.json()["operation_id"]}

    # Las lecturas no se persisten por defecto.
    read = client.get("/api/sheets")
    assert client.get(f"/api/operations/{read.headers['X-Operation-ID']}").status_code == 404


def test_operation_lookup_validates_format(client):
    assert client.get("/api/operations/abc").status_code == 400


def test_logs_include_operation_id(client, caplog):
    import logging

    with caplog.at_level(logging.INFO, logger="claro.api"):
        res = client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "S-1"},
                          headers=OPERATOR)
    op_id = res.headers["X-Operation-ID"]
    records = [r for r in caplog.records if getattr(r, "operation_id", None) == op_id]
    assert any(r.getMessage() == "IPs reservadas" for r in records)
    assert all(r.operator == "Operador Prueba" for r in records)


def test_config_exposes_templates_and_centrales(client):
    cfg = client.get("/api/config").json()
    assert set(cfg["services"]) == {"INTERNET CORPORATIVO", "DATOS", "ACCESO EMPRESARIAL"}
    assert cfg["centrales_source"] == "archivo"
    assert cfg["centrales"][0]["equipos"]
    sync = client.post("/api/centrales/sync", headers=OPERATOR)
    assert sync.status_code == 200
    assert client.get("/api/config").json()["centrales_source"] == "firestore"


def test_service_search(client):
    client.post("/api/reservations", json={"ips": ["10.20.38.6"], "service_id": "S-1"}, headers=OPERATOR)
    client.post("/api/altas", json={"data": {"id_servicio": "S-1", "cliente": "X"}}, headers=OPERATOR)
    found = client.get("/api/services/S-1").json()
    assert [i["ip"] for i in found["ips"]] == ["10.20.38.6"]
    assert len(found["altas"]) == 1
