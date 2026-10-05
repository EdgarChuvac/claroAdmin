"""Regresiones de la revisión de código (normalización, errores, concurrencia, exportación)."""

import io
import re

import openpyxl
from conftest import OPERATOR

OP_RE = re.compile(r"^OP-\d{8}-[0-9A-F]{12}$")


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
        for cells in ws.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"  # texto literal, como lo guardaría Excel con apóstrofo
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client, content, mode="merge"):
    data = {"mode": mode, "confirm_replace": "true"}
    return client.post("/api/inventory/import", files={"file": ("x.xlsx", content)}, data=data, headers=OPERATOR)


SUBNET = [[0, "VLAN 5"], [1, "GW"], [2, ""], [3, ""], [4, ""], [5, ""], [6, ""], [7, "BROADCAST"]]


def test_sheet_names_are_normalized(client):
    res = _import(client, _xlsx({"10.9.9.0 ": SUBNET}))
    assert res.status_code == 200, res.text
    blocks = client.get("/api/blocks", params={"sheet": "10.9.9.0"}).json()["blocks"]
    assert blocks and blocks[0]["sheet_name"] == "10.9.9.0"
    client.post("/api/reservations", json={"ips": ["10.9.9.2"], "service_id": "S-1"}, headers=OPERATOR)
    sheet = next(s for s in client.get("/api/sheets").json()["details"] if s["sheet"] == "10.9.9.0")
    assert sheet["assigned_count"] == 1
    # merge posterior no pisa la reserva
    again = _import(client, _xlsx({"10.9.9.0": SUBNET})).json()
    assert again["conflicts_count"] == 1
    assert client.get("/api/services/S-1").json()["ips"][0]["ip"] == "10.9.9.2"


def test_sheet_must_be_a_24_base(client):
    res = _import(client, _xlsx({"10.9.9.128": SUBNET, "10.9.8.0": SUBNET}))
    assert res.json()["ignored_sheets"] == ["10.9.9.128"]
    dup = _import(client, _xlsx({"10.9.7.0": SUBNET, "10.9.7.0 ": SUBNET}))
    assert dup.status_code == 400


def test_invalid_query_params_return_400(client):
    assert client.get("/api/blocks", params={"sheet": "1.2.3/4"}).status_code == 400
    assert client.get("/api/blocks", params={"sheet": "__abc__"}).status_code in (400, 422)
    assert client.get("/api/services/=abc").status_code == 400
    assert client.get("/api/services/%20").status_code == 400


def test_unhandled_error_keeps_operation_id(client, monkeypatch):
    def boom():
        raise RuntimeError("fallo inesperado")

    monkeypatch.setattr(client.app.state.claro.repo, "list_sheets", boom)
    res = client.get("/api/sheets")
    assert res.status_code == 500
    op_id = res.json()["operation_id"]
    assert OP_RE.match(op_id) and res.headers["X-Operation-ID"] == op_id
    op = client.get(f"/api/operations/{op_id}").json()
    assert op["status_code"] == 500
    assert "fallo inesperado" not in str(op)


def test_404_is_not_persisted(client):
    res = client.get("/api/altas/NO-EXISTE")
    assert res.status_code == 404
    assert client.get(f"/api/operations/{res.json()['operation_id']}").status_code == 404


def test_assigned_first_host_is_not_gateway(client):
    rows = [[0, "VLAN 5"], [1, "SRV-A"], [2, ""], [3, "BROADCAST"]]
    _import(client, _xlsx({"10.9.6.0": rows}))
    block = client.get("/api/blocks", params={"sheet": "10.9.6.0"}).json()["blocks"][0]
    assert block["gateway_ip"] is None
    assert block["assigned_ips"][0]["id"] == "SRV-A"
    exported = client.get("/api/inventory/export").content
    ws = openpyxl.load_workbook(io.BytesIO(exported))["10.9.6.0"]
    assert ws["B2"].value == "SRV-A"


def test_export_never_writes_formulas_and_roundtrips(client):
    rows = [[0, "=HYPERLINK(\"x\")"], [1, "GW"], [2, "+SRV"], [3, "BROADCAST"]]
    _import(client, _xlsx({"10.9.5.0": rows}))
    exported = client.get("/api/inventory/export").content
    ws = openpyxl.load_workbook(io.BytesIO(exported))["10.9.5.0"]
    assert ws["B1"].data_type == "s" and ws["B1"].value == '=HYPERLINK("x")'
    assert ws["B3"].value == "+SRV"
    body = _import(client, exported).json()
    assert body["ips_written"] == 0 and body["conflicts_count"] == 0


def test_resubnet_removes_stale_segments(client):
    _import(client, _xlsx({"10.9.4.0": [[0, "VLAN 1"], [1, "GW"], *[[o, ""] for o in range(2, 15)], [15, "BROADCAST"]]}))
    client.post("/api/reservations", json={"ips": ["10.9.4.3"], "service_id": "S-K"}, headers=OPERATOR)
    two = [[0, "VLAN 1"], [1, "GW"], *[[o, ""] for o in range(2, 7)], [7, "BROADCAST"],
           [8, "VLAN 2"], [9, "GW"], *[[o, ""] for o in range(10, 15)], [15, "BROADCAST"]]
    _import(client, _xlsx({"10.9.4.0": two}))
    blocks = client.get("/api/blocks", params={"sheet": "10.9.4.0"}).json()["blocks"]
    assert [b["network_ip"] for b in blocks] == ["10.9.4.0/29", "10.9.4.8/29"]
    assert [a["ip"] for a in blocks[0]["assigned_ips"]] == ["10.9.4.3"]


def test_guarded_writes_skip_concurrently_modified_ips(repo):
    repo.reserve(["10.20.38.6"], "S-NUEVO", "op", "OP-X")
    # La importación leyó la IP como DISPONIBLE antes de la reserva concurrente.
    stale_expected = {"status": "DISPONIBLE", "service_id": None}
    new_doc = {"ip": "10.20.38.6", "status": "DISPONIBLE", "service_id": None, "segment_id": "10.20.38.0_27",
               "sheet": "10.20.38.0", "octet": 6}
    written, conflicts = repo._guarded_ip_writes([("10.20.38.6", stale_expected, new_doc)])
    assert written == 0 and conflicts[0]["ip"] == "10.20.38.6"
    assert repo.search_service("S-NUEVO")[0]["ip"] == "10.20.38.6"


def test_alta_does_not_store_psk_in_form_data(client):
    res = client.post("/api/altas", json={"data": {"id_servicio": "S-1", "cliente": "X", "psk": "SECRETA",
                                           "sin_factibilidad": True}},
                      headers=OPERATOR)
    doc = client.get(f"/api/altas/{res.json()['alta_id']}").json()
    assert "psk" not in doc["form_data"]


def test_concurrent_identical_altas_create_one(client):
    from concurrent.futures import ThreadPoolExecutor

    payload = {"data": {"id_servicio": "S-C", "cliente": "X", "sin_factibilidad": True, "loopback_auto": True}}
    with ThreadPoolExecutor(max_workers=6) as pool:
        results = list(pool.map(lambda _: client.post("/api/altas", json=payload, headers=OPERATOR).json(), range(6)))
    assert len({r["alta_id"] for r in results}) == 1
    assert sum(not r["duplicate"] for r in results) == 1
