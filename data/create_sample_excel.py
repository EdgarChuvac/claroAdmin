import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def create_sample_excel(filepath: str):
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb = openpyxl.Workbook()
    # Sheet 1: 10.20.38.0
    ws1 = wb.active
    ws1.title = "10.20.38.0"
    
    # Styling colors matching Claro Excel screenshot
    green_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")  # Green for VLAN/Service
    peach_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")  # Peach for GW & client IDs
    grey_fill = PatternFill(start_color="A6A6A6", end_color="A6A6A6", fill_type="solid")   # Grey for Broadcast
    
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    
    # We will populate blocks according to the user's screenshot:
    # Col pair 1 (A & B): 0 to 31 (/27)
    # Col pair 2 (C & D): 32 to 63 (/27)
    # Col pair 3 (E & F): 64 to 95 (/27)
    # Col pair 4 (G & H): 96 to 127 (/27)
    # Col pair 5 (I & J): 128 to 159 (/27)
    # Col pair 6 (K & L): 160 to 191 (/27)
    # Col pair 7 (M & N): 192 to 207 (/28) AND 208 to 223 (/28)
    # Col pair 8 (O & P): 224 to 239 (/28) AND 240 to 255 (/28)
    
    blocks = [
        # (col_num, start_octet, end_octet, vlan_name, row_offset, clients_dict)
        (1, 0, 31, "INTERNET 3740", 1, {
            2: "SERVICIO-DEMO-A",
            3: "SERVICIO-DEMO-B",
            4: "SERVICIO-DEMO-C",
            5: "SERVICIO-DEMO-D"
        }),
        (3, 32, 63, "INTERNET 3740", 1, {}),
        (5, 64, 95, "INTERNET 3740", 1, {}),
        (7, 96, 127, "INTERNET 3740", 1, {}),
        (9, 128, 159, "INTERNET 3740", 1, {}),
        (11, 160, 191, "INTERNET 3740", 1, {}),
        # Stacked /28 blocks in col pair 13 & 14
        (13, 192, 207, "INTERNET 3740", 1, {}),
        (13, 208, 223, "INTERNET 3740", 17, {}),
        # Stacked /28 blocks in col pair 15 & 16
        (15, 224, 239, "INTERNET 3740", 1, {}),
        (15, 240, 255, "INTERNET 3740", 17, {})
    ]
    
    for (col, start_o, end_o, vlan_name, start_row, clients) in blocks:
        for idx, octet in enumerate(range(start_o, end_o + 1)):
            r = start_row + idx
            c_num = ws1.cell(row=r, column=col)
            c_lbl = ws1.cell(row=r, column=col + 1)
            
            c_num.value = octet
            c_num.font = normal_font
            c_num.alignment = Alignment(horizontal="right", vertical="center")
            c_num.border = thin_border
            
            c_lbl.font = normal_font
            c_lbl.border = thin_border
            c_lbl.alignment = Alignment(horizontal="left", vertical="center")
            
            if octet == start_o:
                c_lbl.value = vlan_name
                c_lbl.fill = green_fill
                c_lbl.font = bold_font
            elif octet == start_o + 1:
                c_lbl.value = "GW"
                c_lbl.fill = peach_fill
                c_lbl.font = bold_font
            elif octet == end_o:
                c_lbl.value = "BROADCAST"
                c_lbl.fill = grey_fill
                c_lbl.font = bold_font
            elif octet in clients:
                c_lbl.value = clients[octet]
                c_lbl.fill = peach_fill
            else:
                c_lbl.value = None  # Free / available
                
    # Also add sheet 2: 10.20.37.0
    ws2 = wb.create_sheet(title="10.20.37.0")
    # Quick /27 block for 10.20.37.0
    for idx, octet in enumerate(range(0, 32)):
        r = 1 + idx
        ws2.cell(row=r, column=1, value=octet)
        lbl = ws2.cell(row=r, column=2)
        if octet == 0:
            lbl.value = "INTERNET 3740"
            lbl.fill = green_fill
        elif octet == 1:
            lbl.value = "GW"
            lbl.fill = peach_fill
        elif octet == 31:
            lbl.value = "BROADCAST"
            lbl.fill = grey_fill
        else:
            lbl.value = None

    wb.save(filepath)
    print(f"Sample Excel created at: {filepath}")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "data/ejemplo_inventario_ips.xlsx"
    create_sample_excel(target)
