import math
import re
import os
import tempfile
from ipaddress import IPv4Address, IPv4Network
from typing import Dict, List, Optional, Any
import openpyxl
from openpyxl.styles import PatternFill

class IPBlock:
    def __init__(self, sheet_name: str, network_octet: int, broadcast_octet: int, 
                 col_num: int, start_row: int, end_row: int, vlan_raw: str):
        self.sheet_name = sheet_name
        self.network_octet = network_octet
        self.broadcast_octet = broadcast_octet
        self.col_num = col_num          # Excel 1-based column for octet numbers
        self.val_col_num = col_num + 1  # Excel 1-based column for labels/IDs
        self.start_row = start_row
        self.end_row = end_row
        self.vlan_raw = vlan_raw.strip() if vlan_raw else ""
        
        # Extract VLAN number if possible
        vlan_match = re.search(r'\b(\d{1,5})\b', self.vlan_raw)
        self.vlan = vlan_match.group(1) if vlan_match else ""
        
        if not 0 <= network_octet <= broadcast_octet <= 255:
            raise ValueError("Rango de subred inválido")

        # Calculate CIDR
        self.size = broadcast_octet - network_octet + 1
        if self.size > 0 and (self.size & (self.size - 1)) == 0:
            self.cidr = 32 - int(math.log2(self.size))
        else:
            raise ValueError("El rango de la subred debe ser una potencia de dos")
            
        # Base IP
        # Sheet name usually e.g. "10.20.38.0" -> base is "10.20.38"
        sheet_ip = IPv4Address(sheet_name.strip())
        self.ip_base = ".".join(str(sheet_ip).split(".")[:3])
        network = IPv4Network(f"{self.ip_base}.{network_octet}/{self.cidr}", strict=True)
            
        self.network_ip = str(network)
        self.broadcast_ip = f"{self.ip_base}.{self.broadcast_octet}"
        self.gateway_octet = None
        self.gateway_ip = None
        self.hosts = []        # All hosts in this block
        self.available_ips = [] # List of dicts: {'octet': int, 'ip': str, 'row': int}
        self.assigned_ips = []  # List of dicts: {'octet': int, 'ip': str, 'id': str, 'row': int}

    def to_dict(self) -> Dict[str, Any]:
        return {
            "sheet_name": self.sheet_name,
            "network_ip": self.network_ip,
            "network_octet": self.network_octet,
            "broadcast_ip": self.broadcast_ip,
            "broadcast_octet": self.broadcast_octet,
            "gateway_ip": self.gateway_ip,
            "gateway_octet": self.gateway_octet,
            "vlan": self.vlan,
            "vlan_raw": self.vlan_raw,
            "cidr": f"/{self.cidr}",
            "size": self.size,
            "col_num": self.col_num,
            "val_col_num": self.val_col_num,
            "available_count": len(self.available_ips),
            "assigned_count": len(self.assigned_ips),
            "available_ips": self.available_ips,
            "assigned_ips": self.assigned_ips,
            "first_available": self.available_ips[0] if self.available_ips else None
        }


class ExcelIPAMReader:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.wb = None

    def load_workbook(self, data_only: bool = True):
        self.wb = openpyxl.load_workbook(self.file_path, data_only=data_only)

    def get_sheet_names(self) -> List[str]:
        if not self.wb:
            self.load_workbook()
        return self.wb.sheetnames

    def parse_sheet(self, sheet_name: str) -> List[IPBlock]:
        if not self.wb:
            self.load_workbook()
            
        if sheet_name not in self.wb.sheetnames:
            return []
            
        ws = self.wb[sheet_name]
        blocks: List[IPBlock] = []
        max_col = ws.max_column
        max_row = ws.max_row
        
        # Iterate over columns in pairs (1, 2), (3, 4), etc.
        for c in range(1, max_col + 1, 2):
            val_c = c + 1
            if val_c > max_col:
                break
                
            # Scan down this pair of columns to find blocks
            r = 1
            while r <= max_row:
                num_val = ws.cell(row=r, column=c).value
                label_val = ws.cell(row=r, column=val_c).value
                
                # Check if num_val is an integer octet (0 to 255)
                if num_val is not None:
                    try:
                        octet_num = int(str(num_val).strip())
                    except ValueError:
                        r += 1
                        continue
                        
                    label_str = str(label_val).strip() if label_val is not None else ""
                    
                    if not 0 <= octet_num <= 255:
                        r += 1
                        continue

                    # If this row starts a subnet (not BROADCAST, not GW, and often has VLAN or starts range)
                    # We search until BROADCAST
                    if "BROADCAST" not in label_str.upper():
                        network_octet = octet_num
                        start_row = r
                        vlan_raw = label_str
                        
                        # Find the end of this block (row with BROADCAST)
                        end_r = r + 1
                        found_broadcast = False
                        broadcast_octet = None
                        
                        while end_r <= max_row:
                            end_num_val = ws.cell(row=end_r, column=c).value
                            end_lbl_val = ws.cell(row=end_r, column=val_c).value
                            
                            if end_num_val is None:
                                end_r += 1
                                continue
                                
                            try:
                                curr_octet = int(str(end_num_val).strip())
                            except ValueError:
                                end_r += 1
                                continue
                                
                            curr_lbl = str(end_lbl_val).strip() if end_lbl_val is not None else ""
                            if "BROADCAST" in curr_lbl.upper():
                                broadcast_octet = curr_octet
                                found_broadcast = True
                                break
                            end_r += 1
                            
                        if found_broadcast and broadcast_octet is not None:
                            try:
                                block = IPBlock(
                                    sheet_name=sheet_name,
                                    network_octet=network_octet,
                                    broadcast_octet=broadcast_octet,
                                    col_num=c,
                                    start_row=start_row,
                                    end_row=end_r,
                                    vlan_raw=vlan_raw
                                )
                            except ValueError:
                                r = end_r + 1
                                continue
                            
                            # Parse hosts between start_row and end_r
                            for host_r in range(start_row + 1, end_r):
                                h_num = ws.cell(row=host_r, column=c).value
                                h_lbl = ws.cell(row=host_r, column=val_c).value
                                if h_num is None:
                                    continue
                                try:
                                    h_octet = int(str(h_num).strip())
                                except ValueError:
                                    continue
                                    
                                h_str = str(h_lbl).strip() if h_lbl is not None else ""
                                h_ip = f"{block.ip_base}.{h_octet}"
                                
                                if "GW" in h_str.upper():
                                    block.gateway_octet = h_octet
                                    block.gateway_ip = h_ip
                                elif h_str == "" or h_str == "None" or h_str == "-":
                                    # Available
                                    block.available_ips.append({
                                        "octet": h_octet,
                                        "ip": h_ip,
                                        "row": host_r,
                                        "status": "DISPONIBLE"
                                    })
                                else:
                                    # Assigned to client ID
                                    block.assigned_ips.append({
                                        "octet": h_octet,
                                        "ip": h_ip,
                                        "id": h_str,
                                        "row": host_r,
                                        "status": "OCUPADA"
                                    })
                                    
                            # Default gateway if not explicitly labelled with 'GW'
                            if block.gateway_ip is None:
                                block.gateway_octet = network_octet + 1
                                block.gateway_ip = f"{block.ip_base}.{block.gateway_octet}"
                                
                            blocks.append(block)
                            r = end_r + 1
                            continue
                            
                r += 1
                
        return blocks

    def assign_ip(self, sheet_name: str, ip: str, client_id: str) -> bool:
        """
        Assigns an available IP after resolving its cell from the inventory itself.
        """
        try:
            requested_ip = str(IPv4Address(ip))
        except ValueError:
            return False

        target = None
        for block in self.parse_sheet(sheet_name):
            match = next(
                (item for item in block.available_ips if item["ip"] == requested_ip),
                None,
            )
            if match:
                target = (match["row"], block.val_col_num)
                break
        if target is None:
            return False

        wb_write = openpyxl.load_workbook(self.file_path)
        if sheet_name not in wb_write.sheetnames:
            return False

        ws = wb_write[sheet_name]
        cell = ws.cell(row=target[0], column=target[1])
        current_value = str(cell.value).strip() if cell.value is not None else ""
        if current_value not in ("", "None", "-"):
            return False
        cell.value = client_id
        
        # Orange fill for assigned client (Claro style peach/orange)
        orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
        cell.fill = orange_fill
        
        file_path = os.path.abspath(self.file_path)
        fd, temp_name = tempfile.mkstemp(
            dir=os.path.dirname(file_path), suffix=".xlsx"
        )
        os.close(fd)
        try:
            wb_write.save(temp_name)
            os.replace(temp_name, file_path)
        finally:
            if os.path.exists(temp_name):
                os.unlink(temp_name)
        # Reload internal wb
        self.load_workbook()
        return True
